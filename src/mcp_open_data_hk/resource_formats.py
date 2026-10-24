"""Lossless bounded readers. No guessed units, dates, identifiers or table joins."""

import csv
import io
import json
import math
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime

from defusedxml import ElementTree

from .contracts import DomainError
from .resource_fields import inspect_nested
from .source_contracts import CSV_DATE_PREAMBLE, date_only_header

MAX_ROWS = 100_000
MAX_FIELDS = 1000
MAX_EXPANDED_BYTES = 32_000_000
MAX_ARCHIVE_MEMBERS = 2000
MAX_NORMALIZED_BYTES = 8_000_000
ZIP_SIGNATURES = (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


class ArchiveBudget:
    def __init__(self):
        self.members = 0
        self.expanded = 0

    def member(self):
        self.members += 1
        if self.members > MAX_ARCHIVE_MEMBERS:
            invalid("Archive member budget exceeded", "SOURCE_TOO_LARGE")

    def read(self, archive, member):
        self.member()
        result = bytearray()
        with archive.open(member) as stream:
            bounded = BoundedMember(stream, self)
            while chunk := bounded.read(65536):
                result.extend(chunk)
        return bytes(result)


class BoundedMember:
    def __init__(self, stream, budget):
        self.stream, self.budget = stream, budget

    def read(self, size=-1):
        if size < 0:
            parts = []
            while chunk := self.read(65536):
                parts.append(chunk)
            return b"".join(parts)
        chunk = self.stream.read(
            min(size, MAX_EXPANDED_BYTES - self.budget.expanded + 1)
        )
        self.budget.expanded += len(chunk)
        if self.budget.expanded > MAX_EXPANDED_BYTES:
            invalid("Actual archive expansion exceeds budget", "SOURCE_TOO_LARGE")
        return chunk

    def __getattr__(self, key):
        return getattr(self.stream, key)

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.stream.close()


class BoundedWorkbook(zipfile.ZipFile):
    def __init__(self, body, budget):
        super().__init__(io.BytesIO(body))
        self.budget = budget

    def open(self, name, mode="r", pwd=None, *, force_zip64=False):
        self.budget.member()
        return BoundedMember(
            super().open(name, mode, pwd, force_zip64=force_zip64), self.budget
        )


def normalized_size(value):
    size = 0
    for chunk in json.JSONEncoder(ensure_ascii=False, default=str).iterencode(value):
        size += len(chunk.encode("utf-8"))
        if size > MAX_NORMALIZED_BYTES:
            invalid("Normalized output exceeds budget", "SOURCE_TOO_LARGE")
    return size


def is_archive(body):
    return body.startswith(ZIP_SIGNATURES) or zipfile.is_zipfile(io.BytesIO(body))


def workbook_archive(body, budget):
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        names = set(archive.namelist())
        if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
            invalid("ZIP member is not a supported workbook", "NESTED_ARCHIVE")
        if len(names) > MAX_ARCHIVE_MEMBERS:
            invalid("Workbook member budget exceeded", "SOURCE_TOO_LARGE")
        for member in archive.infolist():
            expanded = budget.read(archive, member)
            if is_archive(expanded):
                invalid("Nested archive inside workbook", "NESTED_ARCHIVE")


def document_limits():
    import resource

    resource.setrlimit(resource.RLIMIT_CPU, (5, 5))
    resource.setrlimit(resource.RLIMIT_FSIZE, (8_000_000, 8_000_000))
    if sys.platform.startswith("linux"):
        resource.setrlimit(resource.RLIMIT_AS, (512_000_000, 512_000_000))


def invalid(message, code="SOURCE_SCHEMA_CHANGED"):
    raise DomainError(code, message, 422)


def text(body):
    for encoding in ("utf-8-sig", "utf-16", "big5hkscs"):
        if encoding == "utf-16" and not body.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        try:
            return body.decode(encoding), encoding
        except UnicodeError:
            continue
    invalid("Unknown text encoding", "UNSUPPORTED_ENCODING")


def at_path(value, path):
    for part in path.split("/") if path else []:
        part = part.replace("~1", "/").replace("~0", "~")
        if isinstance(value, list) and part.isdigit() and int(part) < len(value):
            value = value[int(part)]
        elif isinstance(value, dict) and part in value:
            value = value[part]
        else:
            invalid(
                "Record path is absent; inspect available paths", "INVALID_RECORD_PATH"
            )
    return value


def paths(value, prefix="", depth=0):
    if depth > 8:
        return []
    if isinstance(value, list):
        return [prefix] + (
            paths(value[0], prefix + "/0", depth + 1)
            if len(value) == 1 and isinstance(value[0], dict)
            else []
        )
    if isinstance(value, dict):
        return [
            p
            for k, v in value.items()
            for p in paths(v, prefix + ("/" if prefix else "") + k, depth + 1)
        ]
    return []


def safe_json(value, depth=0):
    if depth > 30:
        invalid("Nested data exceeds depth budget", "SOURCE_TOO_COMPLEX")
    if isinstance(value, float) and not math.isfinite(value):
        invalid("Nonfinite numeric value")
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, dict):
        return {str(k): safe_json(v, depth + 1) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [safe_json(v, depth + 1) for v in value]
    return value


def no_duplicates(pairs):
    result = {}
    for k, v in pairs:
        if k in result:
            invalid("Duplicate JSON object key")
        result[k] = v
    return result


def table(matrix, info=None, positional_date_header=False):
    matrix = iter(matrix)
    header = next(matrix, None)
    if header is None:
        invalid("No table header")
    raw = [str(x) if x is not None else "" for x in header]
    dated = positional_date_header and date_only_header(raw)
    if not raw or len(raw) > MAX_FIELDS:
        invalid("Table field budget exceeded", "SOURCE_TOO_COMPLEX")
    header = []
    mapping = []
    for i, name in enumerate(raw):
        field = (
            name
            if not dated and name and raw.count(name) == 1
            else "_column_" + str(i + 1)
        )
        while field in header or (field in raw and field != name):
            field = "_" + field
        header.append(field)
        mapping.append({"name": field, "sourceHeader": name, "columnIndex": i + 1})
    rows = []
    normalized = 2
    for values in matrix:
        if len(rows) >= MAX_ROWS:
            invalid("Source exceeds row budget", "SOURCE_TOO_LARGE")
        if len(values) > MAX_FIELDS:
            invalid("Table field budget exceeded", "SOURCE_TOO_COMPLEX")
        if all(x is None or x == "" for x in values):
            continue
        while len(values) > len(header):
            i = len(header)
            field = "_column_" + str(i + 1)
            while field in header:
                field = "_" + field
            header.append(field)
            mapping.append({"name": field, "sourceHeader": None, "columnIndex": i + 1})
        row = dict(zip(header, values))
        normalized += normalized_size(row) + 2
        if normalized > MAX_NORMALIZED_BYTES:
            invalid("Normalized table exceeds budget", "SOURCE_TOO_LARGE")
        rows.append(row)
    if info is not None:
        if dated:
            info["readerContract"] = CSV_DATE_PREAMBLE
            info["rowMeaning"] = (
                "The dated first row is retained as header metadata; remaining raw rows may include source header/description rows. Row counts are not verified counts of real-world entities."
            )
        info["sourceColumns"] = mapping
        info["headerVerification"] = (
            "First row retained as source header metadata; blank/duplicate/extra columns have explicit positional names. No semantic table interpretation."
        )
    return rows


def parse(body, fmt, record_path="", sheet="", *, _budget=None, _member=False):
    fmt = fmt.upper()
    budget = _budget or ArchiveBudget()
    archived = is_archive(body)
    if archived and _member and fmt not in {"XLSX", "XLSM"}:
        invalid("Nested record archives are not supported", "NESTED_ARCHIVE")
    if archived and fmt not in {"XLSX", "XLSM"}:
        with zipfile.ZipFile(io.BytesIO(body)) as archive:
            members = [m for m in archive.infolist() if not m.is_dir()]
            if (
                len(archive.infolist()) > 1000
                or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES
            ):
                invalid("Archive expansion exceeds budget", "SOURCE_TOO_LARGE")
            allowed = {
                "CSV",
                "JSON",
                "GEOJSON",
                "XML",
                "RSS",
                "GML",
                "KML",
                "XLS",
                "XLSX",
                "XLSM",
                "PDF",
            }
            selected = [
                m for m in members if m.filename.rsplit(".", 1)[-1].upper() in allowed
            ]
            if not selected:
                invalid(
                    "Archive contains no supported record files", "UNSUPPORTED_FORMAT"
                )
            if any(m.flag_bits & 1 for m in selected):
                invalid("Encrypted archive", "ACCESS_RESTRICTED")
            all_rows = []
            member_info = []
            normalized = 2
            for m in archive.infolist():
                member_body = budget.read(archive, m)
                child_format = m.filename.rsplit(".", 1)[-1].upper()
                # Inspect even unread or empty/disguised members before any recursion.
                if is_archive(member_body) and child_format not in {"XLSX", "XLSM"}:
                    invalid(
                        "Nested record archives are not supported", "NESTED_ARCHIVE"
                    )
                if m not in selected:
                    continue
                child_format = m.filename.rsplit(".", 1)[-1].upper()
                child_rows, child_schema = parse(
                    member_body,
                    child_format,
                    record_path,
                    sheet,
                    _budget=budget,
                    _member=True,
                )
                member_info.append(
                    {
                        "name": m.filename,
                        "format": child_format,
                        "recordCount": len(child_rows),
                    }
                )
                for row in child_rows:
                    if "_sourceArchiveMember" in row:
                        invalid("Reserved archive provenance field collision")
                    record = {**row, "_sourceArchiveMember": m.filename}
                    normalized += normalized_size(record) + 2
                    if normalized > MAX_NORMALIZED_BYTES:
                        invalid("Normalized archive exceeds budget", "SOURCE_TOO_LARGE")
                    all_rows.append(record)
                if len(all_rows) > MAX_ROWS:
                    invalid("Archive exceeds row budget", "SOURCE_TOO_LARGE")
            rows, info = parse(
                json.dumps(all_rows, ensure_ascii=False).encode(), "JSON"
            )
            info.update(
                format=fmt,
                scope="supported_archive_members",
                archiveMembers=member_info,
                unreadMembers=[m.filename for m in members if m not in selected],
                datasetCompleteness="Only the listed supported archive members; binary/other members remain unqueried.",
            )
            info["queryFields"] = inspect_nested(rows)
            return rows, info
    info = {
        "format": fmt,
        "recordPath": record_path or None,
        "sheet": sheet or None,
        "normalization": "Source values retained; no inferred units or reference periods.",
    }
    if fmt == "CSV":
        decoded, encoding = text(body)
        info["encoding"] = encoding
        if "\x00" in decoded or decoded.lstrip().lower().startswith(
            ("<!doctype html", "<html")
        ):
            invalid("Source returned binary or HTML rather than CSV")
        sample = decoded[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = table(
            csv.reader(io.StringIO(decoded, newline=""), dialect, strict=True),
            info,
            positional_date_header=True,
        )
    elif fmt in {"JSON", "GEOJSON", "API"}:
        decoded, encoding = text(body)
        value = json.loads(
            decoded,
            object_pairs_hook=no_duplicates,
            parse_constant=lambda _: invalid("Nonfinite JSON number"),
        )
        info["encoding"] = encoding
        candidates = paths(value)
        info["availableRecordPaths"] = candidates[:100]

        def context(v, depth=0):
            if depth > 8:
                return {"omitted": "depth limit"}
            if isinstance(v, list):
                return {"sourceArrayLength": len(v)}
            if isinstance(v, dict):
                return {k: context(x, depth + 1) for k, x in v.items()}
            return v

        if isinstance(value, dict):
            info["sourceContext"] = context(value)
        if isinstance(value, dict) and (
            value.get("success") is False
            or isinstance(value.get("error"), dict)
            and "code" in value["error"]
        ):
            invalid("Source API returned an error envelope", "SOURCE_API_ERROR")
        if (
            isinstance(value, dict)
            and set(value) == {"message"}
            and str(value["message"]).strip().casefold()
            in {"no record found", "no records found", "no data found", "no data"}
        ):
            info["sourceContext"] = dict(value)
            info["sourceResult"] = "empty"
            value = []
        calendar = value.get("vcalendar") if isinstance(value, dict) else None
        calendar = (
            calendar[0]
            if isinstance(calendar, list)
            and len(calendar) == 1
            and isinstance(calendar[0], dict)
            and isinstance(calendar[0].get("vevent"), list)
            and "prodid" in calendar[0]
            and "version" in calendar[0]
            else None
        )
        if calendar is not None:
            info["calendarMetadata"] = {
                k: v for k, v in calendar.items() if k != "vevent"
            }
        if calendar is not None and record_path in {"", "vcalendar/0/vevent"}:
            info["recordPath"] = "vcalendar/0/vevent"
            info["scope"] = "calendar_events"
            info["rowMeaning"] = (
                "One source iCalendar event; calendar metadata retained separately."
            )
            value = calendar["vevent"]
        elif isinstance(value, dict) and value.get("type") == "FeatureCollection":
            info["sourceCrs"] = value.get("crs")
            value = value.get("features")
            info["recordPath"] = "features"
            info["coordinateReferenceSystem"] = (
                "crs is preserved as source metadata; absent crs follows GeoJSON specification"
            )
        elif record_path:
            value = at_path(value, record_path)
        elif (
            not isinstance(value, list)
            and len([p for p in candidates if "/0/" not in p]) == 1
        ):
            chosen = next(p for p in candidates if "/0/" not in p)
            info["recordPath"] = chosen
            value = at_path(value, chosen)
        elif not isinstance(value, list):
            # Keep a nested API object intact; choosing a nested array is explicit.
            value = [value]
            info["scope"] = "source_response_object"
        if not isinstance(value, list):
            invalid("Record path must select an array", "INVALID_RECORD_PATH")
        rows = [v if isinstance(v, dict) else {"value": v} for v in value]
    elif fmt in {"XLSX", "XLSM"}:
        from openpyxl.reader.excel import ExcelReader

        workbook_archive(body, budget)
        reader = ExcelReader(
            io.BytesIO(body), read_only=True, data_only=False, keep_links=False
        )
        reader.archive.close()
        reader.archive = BoundedWorkbook(body, budget)
        try:
            reader.read()
        except BaseException:
            reader.archive.close()
            raise
        book = reader.wb
        try:
            info["availableSheets"] = book.sheetnames
            if not sheet and len(book.sheetnames) != 1:
                invalid(
                    "Select a workbook sheet: " + ", ".join(book.sheetnames)[:500],
                    "SHEET_REQUIRED",
                )
            if sheet and sheet not in book.sheetnames:
                invalid("Workbook sheet not found", "INVALID_SHEET")
            ws = book[sheet or book.sheetnames[0]]
            info["sheet"] = ws.title
            # Preserve formulas rather than presenting unverified cached results.
            rows = table(ws.values, info)
        finally:
            book.close()
    elif fmt == "XLS":
        import xlrd

        book = xlrd.open_workbook(file_contents=body, on_demand=True)
        try:
            info["availableSheets"] = book.sheet_names()
            if not sheet and book.nsheets != 1:
                invalid(
                    "Select a workbook sheet: " + ", ".join(book.sheet_names())[:500],
                    "SHEET_REQUIRED",
                )
            ws = book.sheet_by_name(sheet) if sheet else book.sheet_by_index(0)
            info["sheet"] = ws.name
            info["warning"] = (
                "Legacy XLS exposes stored cell results; original workbook remains authoritative."
            )
            rows = table((ws.row_values(i) for i in range(ws.nrows)), info)
        finally:
            book.release_resources()
    elif fmt in {"XML", "RSS", "GML", "KML"}:
        root = ElementTree.fromstring(body)
        info["rootTag"] = root.tag
        if fmt == "RSS":
            local = root.tag.rsplit("}", 1)[-1]
            valid = (
                (
                    local == "rss"
                    and any(e.tag.rsplit("}", 1)[-1] == "channel" for e in root)
                )
                or root.tag == "{http://www.w3.org/2005/Atom}feed"
                or (
                    root.tag == "{http://www.w3.org/1999/02/22-rdf-syntax-ns#}RDF"
                    and any(
                        e.tag
                        in {
                            "{http://purl.org/rss/1.0/}channel",
                            "{http://purl.org/rss/1.0/}item",
                        }
                        for e in root
                    )
                )
            )
            if not valid:
                invalid("Published RSS source returned a non-feed document")
            info["readerContract"] = "rss-feed-root-1"
            info["fieldCompatibility"] = (
                "RSS/Atom fields are optional and extensible; observed types are retained, not a fixed analytical schema."
            )
            info["rowMeaning"] = (
                "Raw XML elements under the selected path; default selection is the direct feed container, not a count of all catalogue books."
            )

        def node(e):
            if not len(e) and not e.attrib:
                return e.text or ""
            result = {"@" + k: v for k, v in e.attrib.items()}
            for child in e:
                key = child.tag
                if key in result:
                    if not isinstance(result[key], list):
                        result[key] = [result[key]]
                    result[key].append(node(child))
                else:
                    result[key] = node(child)
            if e.text and e.text.strip():
                result["#text"] = e.text
            return result

        if record_path:
            if ".." in record_path or len(record_path) > 300:
                invalid("Use a simple XML element path", "INVALID_RECORD_PATH")
            elements = root.findall(record_path)
        else:
            elements = list(root)
            if len({e.tag for e in elements}) > 1:
                elements = [root]
                info["scope"] = "source_response_object"
        rows = [v if isinstance(v := node(e), dict) else {e.tag: v} for e in elements]
    elif fmt == "PDF":
        with tempfile.TemporaryFile() as output:
            try:
                subprocess.run(
                    [sys.executable, "-m", "mcp_open_data_hk.document_reader"],
                    input=body,
                    stdout=output,
                    stderr=subprocess.DEVNULL,
                    timeout=10,
                    check=True,
                )
                output.seek(0)
                extracted = output.read(8_000_001)
                if len(extracted) > 8_000_000:
                    invalid("Extracted text exceeds budget", "SOURCE_TOO_LARGE")
                rows = json.loads(extracted)
            except (subprocess.SubprocessError, ValueError):
                invalid(
                    "PDF extraction failed or exceeded its isolated resource budget",
                    "DOCUMENT_EXTRACTION_FAILED",
                )
        info.update(
            scope="document_pages",
            verification="unverified_text_extraction",
            structuredQueryable=False,
            emptyPages=[r["page"] for r in rows if not r["text"].strip()],
        )
    else:
        invalid("No reader for this source format", "UNSUPPORTED_FORMAT")
    if len(rows) > MAX_ROWS:
        invalid("Source exceeds row budget", "SOURCE_TOO_LARGE")
    if sum(len(r) for r in rows) > 2_000_000:
        invalid("Source exceeds cell budget", "SOURCE_TOO_COMPLEX")
    normalized_size(rows)
    rows = safe_json(rows)
    fields = sorted({k for r in rows for k in r})
    if len(fields) > MAX_FIELDS:
        invalid("Source exceeds field budget", "SOURCE_TOO_COMPLEX")

    def types(v):
        if v is None:
            return "null"
        if isinstance(v, bool):
            return "boolean"
        if isinstance(v, (int, float)):
            return "number"
        if isinstance(v, dict):
            return "object"
        if isinstance(v, list):
            return "array"
        return "string"

    info["fields"] = [
        {
            "name": k,
            "types": sorted({types(r.get(k)) for r in rows}),
            "missingCount": sum(k not in r or r[k] is None for r in rows),
            "unit": None,
        }
        for k in fields
    ]
    info["queryFields"] = inspect_nested(rows)
    return rows, info
